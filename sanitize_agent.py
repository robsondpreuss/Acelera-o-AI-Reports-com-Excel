"""
Porsche Sales Sanitization Agent
--------------------------------
Reads BASE_DADOS.xlsx and applies the rules from schema.md, producing a
new xlsx file with sanitized columns inserted immediately after each
source column. Original columns are preserved.

Usage:
    python sanitize_agent.py <input_xlsx> <output_xlsx>
"""

from __future__ import annotations

import calendar
import re
import sys
from datetime import datetime, date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


# ---------------------------------------------------------------------------
# Canonical reference data
# ---------------------------------------------------------------------------

CANONICAL_MODELS = [
    "911 Carrera", "911 Carrera S", "911 Carrera GTS", "911 Turbo", "911 Turbo S",
    "911 GT3", "911 GT3 RS", "911 Dakar", "911 Targa 4", "911 Targa 4S",
    "718 Cayman", "718 Cayman S", "718 Cayman GT4 RS",
    "718 Boxster", "718 Boxster GTS", "718 Spyder RS",
    "Cayenne", "Cayenne S", "Cayenne Coupe", "Cayenne E-Hybrid",
    "Cayenne Turbo", "Cayenne Turbo GT",
    "Macan", "Macan S", "Macan T", "Macan GTS", "Macan Electric",
    "Panamera", "Panamera 4", "Panamera 4S", "Panamera Turbo",
    "Panamera Turbo S", "Panamera 4 E-Hybrid",
    "Taycan", "Taycan 4S", "Taycan GTS", "Taycan Turbo", "Taycan Turbo S",
    "Taycan Cross Turismo",
]
MODEL_LOOKUP = {re.sub(r"\s+", " ", m).lower(): m for m in CANONICAL_MODELS}

PAYMENT_CANONICAL = {
    "credit card": "Credit Card",
    "creditcard": "Credit Card",
    "credit": "Credit Card",
    "credit card payment": "Credit Card",
    "debit card": "Debit Card",
    "debitcard": "Debit Card",
    "bank transfer": "Bank Transfer",
    "banktransfer": "Bank Transfer",
    "wire transfer": "Wire Transfer",
    "wiretransfer": "Wire Transfer",
    "wire": "Wire Transfer",
    "bank wire": "Wire Transfer",
    "financing": "Financing",
    "financing plan": "Financing",
    "finance": "Financing",
    "lease": "Lease",
    "leasing": "Lease",
    "lease plan": "Lease",
    "cash": "Cash",
    "cash payment": "Cash",
    "ach": "ACH Payment",
    "ach payment": "ACH Payment",
    "crypto": "Crypto Payment",
    "crypto payment": "Crypto Payment",
    "cryptocurrency": "Crypto Payment",
}

DELIVERY_CANONICAL = {
    "delivered": "Delivered",
    "deliverd": "Delivered",  # typo per schema
    "pending": "Pending",
    "in transit": "In Transit",
    "intransit": "In Transit",
    "cancelled": "Cancelled",
    "canceled": "Cancelled",
    "awaiting delivery": "Awaiting Delivery",
    "awaiting pickup": "Awaiting Pickup",
    "pending approval": "Pending Approval",
    "pending review": "Pending Review",
    "shipped": "Shipped",
    "awaiting review": "Awaiting Review",
}

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR",
    "california": "CA", "colorado": "CO", "connecticut": "CT", "delaware": "DE",
    "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID",
    "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
    "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN",
    "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE",
    "nevada": "NV", "new hampshire": "NH", "new jersey": "NJ",
    "new mexico": "NM", "new york": "NY", "north carolina": "NC",
    "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR",
    "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC",
    "south dakota": "SD", "tennessee": "TN", "texas": "TX", "utah": "UT",
    "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
    "district of columbia": "DC",
}
US_STATE_CODES = set(US_STATES.values())

MONTH_NAMES = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTH_ABBR = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}

NUMBER_WORDS = {
    "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
    "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11,
    "twelve": 12, "thirteen": 13, "fourteen": 14, "fifteen": 15,
    "sixteen": 16, "seventeen": 17, "eighteen": 18, "nineteen": 19,
    "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60,
    "seventy": 70, "eighty": 80, "ninety": 90,
}
SCALE_WORDS = {"hundred": 100, "thousand": 1_000, "million": 1_000_000}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def words_to_int(text: str) -> int | None:
    """Convert English number words to int. Returns None if not recognizable."""
    text = text.lower().replace("-", " ").replace(" and ", " ")
    tokens = [t for t in re.split(r"\s+", text) if t]
    if not tokens:
        return None
    total = 0
    current = 0
    for tok in tokens:
        if tok in NUMBER_WORDS:
            current += NUMBER_WORDS[tok]
        elif tok in SCALE_WORDS:
            scale = SCALE_WORDS[tok]
            current = max(current, 1) * scale
            total += current
            current = 0
        else:
            return None
    return total + current


def valid_date(y: int, m: int, d: int) -> bool:
    try:
        date(y, m, d)
        return True
    except ValueError:
        return False


# ---------------------------------------------------------------------------
# Sanitizers
# ---------------------------------------------------------------------------

def sanitize_date(raw) -> str:
    if raw is None or raw == "":
        return "INVALID"
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%d")
    if isinstance(raw, date):
        return raw.strftime("%Y-%m-%d")
    s = str(raw).strip()

    # 1. YYYY-MM-DD / YYYY/MM/DD / YYYY.MM.DD
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f"{y:04d}-{mo:02d}-{d:02d}" if valid_date(y, mo, d) else "INVALID"

    # 2. MM/DD/YYYY or MM-DD-YYYY
    m = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", s)
    if m:
        mo, d, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f"{y:04d}-{mo:02d}-{d:02d}" if valid_date(y, mo, d) else "INVALID"

    # 3. MM/DD/YY or MM-DD-YY (assume 2000-2099)
    m = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{2})", s)
    if m:
        mo, d, yy = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y = 2000 + yy
        return f"{y:04d}-{mo:02d}-{d:02d}" if valid_date(y, mo, d) else "INVALID"

    # 4. Month/Mon DDth[,] YYYY
    m = re.fullmatch(
        r"([A-Za-z]+)\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s
    )
    if m:
        mon_raw = m.group(1).lower()
        d = int(m.group(2))
        y = int(m.group(3))
        mo = MONTH_NAMES.get(mon_raw) or MONTH_ABBR.get(mon_raw)
        if mo is None:
            return "INVALID"
        return f"{y:04d}-{mo:02d}-{d:02d}" if valid_date(y, mo, d) else "INVALID"

    return "INVALID"


def sanitize_model(raw) -> str:
    if raw is None:
        return "INVALID"
    s = re.sub(r"\s+", " ", str(raw)).strip()
    if not s:
        return "INVALID"
    key = s.lower()
    if key in MODEL_LOOKUP:
        return MODEL_LOOKUP[key]
    # Unknown: title-case but preserve common acronyms
    def fix(tok: str) -> str:
        up = tok.upper()
        if up in {"GT3", "GT4", "RS", "GTS", "S", "T", "4S", "4"}:
            return up
        if up.startswith("E-") or "-" in tok:
            return "-".join(p.capitalize() for p in tok.split("-"))
        return tok.capitalize()
    return " ".join(fix(t) for t in s.split())


def sanitize_year(raw) -> str:
    if raw is None or raw == "":
        return "INVALID"
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        y = int(raw)
        return str(y) if 1990 <= y <= 2035 else "INVALID"
    s = str(raw).strip().lower()

    # Direct 4-digit year
    m = re.fullmatch(r"(\d{4})", s)
    if m:
        y = int(m.group(1))
        return str(y) if 1990 <= y <= 2035 else "INVALID"

    # "20-24" or "20 24" -> 2024
    m = re.fullmatch(r"(\d{2})[\s\-](\d{2})", s)
    if m:
        y = int(m.group(1) + m.group(2))
        return str(y) if 1990 <= y <= 2035 else "INVALID"

    # "two thousand twenty four" -> 2024
    n = words_to_int(s)
    if n is not None and 1990 <= n <= 2035:
        return str(n)

    # Colloquial "twenty twenty four" form: split into two halves
    tokens = re.split(r"[\s\-]+", s)
    if len(tokens) >= 2 and tokens[0] in NUMBER_WORDS:
        first = NUMBER_WORDS[tokens[0]]
        rest = words_to_int(" ".join(tokens[1:]))
        if rest is not None and 0 <= rest < 100:
            y = first * 100 + rest
            if 1990 <= y <= 2035:
                return str(y)
    return "INVALID"


def sanitize_price(raw) -> str:
    if raw is None or raw == "":
        return "INVALID"
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return f"{float(raw):.2f}"
    s = str(raw).strip().lower()
    s = s.replace("usd", "").replace("dollars", "").replace("dollar", "")
    s = s.replace("$", "").strip()

    # Detect "k" multiplier
    k_mult = False
    if re.search(r"\d\s*k\b", s):
        k_mult = True
        s = re.sub(r"k\b", "", s)

    s = s.strip()

    # Try word-form ("eighty two thousand")
    if re.fullmatch(r"[a-z\s\-]+", s):
        n = words_to_int(s)
        if n is not None:
            return f"{float(n):.2f}"
        return "INVALID"

    # Numeric cleanup: handle European "." thousand separator and "," decimal.
    # Strategy:
    #  - If last separator is ',' followed by 1-2 digits -> decimal comma
    #  - Else strip dots and commas as thousands separators (except a trailing
    #    '.NN' decimal point with 1-2 digits)
    digits_only = re.sub(r"[^0-9.,]", "", s)
    if not digits_only:
        return "INVALID"

    # Detect decimal: last '.' or ',' followed by exactly 1-2 digits at end
    dec_match = re.search(r"[.,](\d{1,2})$", digits_only)
    if dec_match:
        sep = digits_only[dec_match.start()]
        int_part = digits_only[: dec_match.start()]
        dec_part = dec_match.group(1)
        int_part = re.sub(r"[.,]", "", int_part)
        value = float(f"{int_part}.{dec_part}") if int_part else float(f"0.{dec_part}")
    else:
        value = float(re.sub(r"[.,]", "", digits_only))

    if k_mult:
        value *= 1000
    return f"{value:.2f}"


def sanitize_mileage(raw) -> str:
    if raw is None or raw == "":
        return "0"
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return str(int(round(float(raw))))
    s = str(raw).strip().lower()

    if not s:
        return "0"
    # Zero indicators
    if re.search(r"\b(new|new car|zero(\s+miles)?|0\s*(mi|miles)?)\b", s) and not re.search(r"[1-9]", s):
        return "0"

    is_km = "km" in s

    # Word forms
    word_part = re.sub(r"\b(miles?|mi|km)\b", "", s)
    word_part = re.sub(r"[^a-z\s\-]", "", word_part).strip()
    if word_part and re.fullmatch(r"[a-z\s\-]+", word_part):
        n = words_to_int(word_part)
        if n is not None:
            value = n
            if is_km:
                value = round(value * 0.621371)
            return str(int(value))

    digits_only = re.sub(r"[^0-9.,]", "", s)
    if not digits_only:
        return "0"
    # Treat '.' and ',' as thousand separators (mileage is integer miles/km)
    cleaned = re.sub(r"[.,]", "", digits_only)
    try:
        value = int(cleaned)
    except ValueError:
        return "0"
    if is_km:
        value = round(value * 0.621371)
    return str(value)


def sanitize_payment(raw) -> str:
    if raw is None:
        return "INVALID"
    s = re.sub(r"[_\-]", " ", str(raw)).strip().lower()
    s = re.sub(r"\s+", " ", s)
    if s in PAYMENT_CANONICAL:
        return PAYMENT_CANONICAL[s]
    # Title-case fallback
    return " ".join(w.capitalize() for w in s.split()) if s else "INVALID"


def sanitize_city(raw) -> str:
    if raw is None:
        return "INVALID"
    s = re.sub(r"\s+", " ", str(raw)).strip()
    if not s:
        return "INVALID"
    parts = []
    for token in s.split(" "):
        # Preserve abbreviations like "St." -> "St."
        if re.fullmatch(r"[A-Za-z]+\.", token):
            parts.append(token[0].upper() + token[1:].lower())
        else:
            parts.append(token.capitalize())
    return " ".join(parts)


def sanitize_state(raw) -> str:
    if raw is None:
        return "INVALID"
    s = str(raw).strip()
    if not s:
        return "INVALID"
    upper = s.upper()
    if len(upper) == 2 and upper in US_STATE_CODES:
        return upper
    key = re.sub(r"\s+", " ", s.lower())
    if key in US_STATES:
        return US_STATES[key]
    return "INVALID"


def sanitize_delivery(raw) -> str:
    if raw is None:
        return "INVALID"
    s = str(raw).strip().lower()
    s = re.sub(r"[^a-z\s]", " ", s)  # drop punctuation
    s = re.sub(r"\s+", " ", s).strip()
    if s in DELIVERY_CANONICAL:
        return DELIVERY_CANONICAL[s]
    # Try compressed form for "intransit"
    if s.replace(" ", "") in DELIVERY_CANONICAL:
        return DELIVERY_CANONICAL[s.replace(" ", "")]
    return " ".join(w.capitalize() for w in s.split()) if s else "INVALID"


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

SANITIZERS = {
    "sale_date": ("SaleDateSanitized", sanitize_date),
    "porsche_model": ("PorscheModelSanitized", sanitize_model),
    "model_year": ("ModelYearSanitized", sanitize_year),
    "sale_price": ("SalesPriceSanitized", sanitize_price),
    "vehicle_mileage": ("VehicleMileageSanitized", sanitize_mileage),
    "payment_method": ("PayMethodSanitized", sanitize_payment),
    "city": ("CitySanitized", sanitize_city),
    "state": ("StateSanitized", sanitize_state),
    "delivery_status": ("DeliveryStatusSanitized", sanitize_delivery),
}


def process(input_path: Path, output_path: Path) -> dict:
    wb_in = openpyxl.load_workbook(input_path)
    ws_in = wb_in.active

    headers = [c.value for c in ws_in[1]]
    rows = list(ws_in.iter_rows(min_row=2, values_only=True))

    # Build new header order: each original column, sanitized column right after
    out_headers: list[str] = []
    col_plan: list[tuple[str, str | None]] = []  # (source_col, sanitized_name)
    for h in headers:
        out_headers.append(str(h))
        col_plan.append((h, None))
        if h in SANITIZERS:
            san_name = SANITIZERS[h][0]
            out_headers.append(san_name)
            col_plan.append((h, san_name))

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "sanitized"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    san_fill = PatternFill("solid", fgColor="2E7D32")
    for idx, h in enumerate(out_headers, start=1):
        cell = ws_out.cell(row=1, column=idx, value=h)
        cell.font = header_font
        cell.fill = san_fill if h.endswith("Sanitized") else header_fill

    stats = {"rows": 0, "invalid": {san: 0 for _, (san, _) in SANITIZERS.items()}}
    source_idx = {h: i for i, h in enumerate(headers)}

    for row in rows:
        out_row: list = []
        for source_col, san_name in col_plan:
            raw = row[source_idx[source_col]] if source_col in source_idx else None
            if san_name is None:
                out_row.append(raw)
            else:
                fn = SANITIZERS[source_col][1]
                value = fn(raw)
                out_row.append(value)
                if value == "INVALID":
                    stats["invalid"][san_name] += 1
        ws_out.append(out_row)
        stats["rows"] += 1

    # Auto width
    for col in ws_out.columns:
        letter = col[0].column_letter
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws_out.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)
    ws_out.freeze_panes = "A2"

    wb_out.save(output_path)
    return stats


def main() -> int:
    if len(sys.argv) >= 3:
        in_path = Path(sys.argv[1])
        out_path = Path(sys.argv[2])
    else:
        in_path = Path("/sessions/dazzling-youthful-goldberg/mnt/uploads/BASE_DADOS.xlsx")
        out_path = Path("/sessions/dazzling-youthful-goldberg/mnt/outputs/BASE_DADOS_TRATADA.xlsx")

    stats = process(in_path, out_path)
    print(f"Rows processed: {stats['rows']}")
    print("INVALID counts per sanitized column:")
    for k, v in stats["invalid"].items():
        print(f"  {k}: {v}")
    print(f"Saved to: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
.xlsx")
        out_path = Path("/sessions/dazzling-youthful-goldberg/mnt/outputs/BASE_DADOS_TRATADA.xlsx")

    stats = process(in_path, out_path)
    print(f"Rows processed: {stats['rows']}")
    print("INVALID counts per sanitized column:")
    for k, v in stats["invalid"].items():
        print(f"  {k}: {v}")
    print(f"Saved to: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
